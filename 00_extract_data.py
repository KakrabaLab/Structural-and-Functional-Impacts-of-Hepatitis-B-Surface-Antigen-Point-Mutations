"""
Step 0: Data extraction and preparation.

Extracts the raw MD RMSD trajectories from the supplied Excel workbook
(MD_Simulations_1_.xlsx) and builds the graph-theoretic / molar-mass
descriptor tables transcribed from the thesis (Tables 5 and 6 in the
original thesis document). Produces the cleaned CSV/pickle files consumed
by every downstream script (01-06).

Input:
  /mnt/user-data/uploads/MD_Simulations_1_.xlsx
    - Sheet '200 ns Simulation': single shared time column + RMSD per variant
    - Sheet '50 ns Simulation' : one (Time, RMSD) column-pair per variant,
      ragged (different sampling intervals/lengths per variant)

Output (written to analysis/data/):
  rmsd_200ns.csv                    - tidy 200 ns RMSD table (Time_ns + 6 variants)
  rmsd_50ns.pkl                     - dict of {variant: DataFrame(Time_ns, RMSD)} for 50 ns
  table5_graph_descriptors.csv      - d1-d8 weighted graph descriptors per variant
  table6_molar_mass_descriptors.csv - 15-node molar-mass deviation profile per variant
  md_summary_thesis_reported.csv    - thesis-reported 50/200 ns mean & max RMSD (for cross-check)
"""
import openpyxl, pandas as pd, numpy as np, pickle, os

RAW_XLSX = "/mnt/user-data/uploads/MD_Simulations_1_.xlsx"
OUT_DIR = "analysis/data"
os.makedirs(OUT_DIR, exist_ok=True)

VARIANTS = ["WT", "Y100C", "C124Y", "S132F", "M133K", "G145R"]

# ---------------------------------------------------------------------
# 1. 200 ns RMSD trajectories (single shared time column)
# ---------------------------------------------------------------------
wb = openpyxl.load_workbook(RAW_XLSX, data_only=True)

ws = wb["200 ns Simulation"]
rows = list(ws.iter_rows(min_row=2, values_only=True))
df200 = pd.DataFrame(rows, columns=["Time_ns", "WT", "Y100C", "C124Y", "S132F", "M133K", "G145R"])
df200 = df200.dropna(how="all").astype(float)
df200.to_csv(f"{OUT_DIR}/rmsd_200ns.csv", index=False)
print("200 ns:", df200.shape)

# ---------------------------------------------------------------------
# 2. 50 ns RMSD trajectories (ragged, one Time/RMSD column-pair per variant)
# ---------------------------------------------------------------------
ws2 = wb["50 ns Simulation"]
rows2 = list(ws2.iter_rows(min_row=3, values_only=True))
data = {}
for i, v in enumerate(VARIANTS):
    tcol, rcol = i * 3, i * 3 + 1
    t = pd.Series([r[tcol] for r in rows2], dtype="float64")
    r = pd.Series([r[rcol] for r in rows2], dtype="float64")
    mask = t.notna() & r.notna()
    data[v] = pd.DataFrame({"Time_ns": t[mask].values, "RMSD": r[mask].values})
    print(v, data[v].shape, "max_t=", data[v]["Time_ns"].max())

with open(f"{OUT_DIR}/rmsd_50ns.pkl", "wb") as f:
    pickle.dump(data, f)

# ---------------------------------------------------------------------
# 3. Table 5 (thesis) - weighted principal-domain graph descriptors (d1-d8)
#    Transcribed from the thesis manuscript (Chapter 4, "Network Analysis
#    of HBsAg and Its Mutants").
# ---------------------------------------------------------------------
table5 = pd.DataFrame({
    "Variant": ["Wildtype", "Y100C", "C124Y", "S132F", "M133K", "G145R"],
    "d1_density":            [0.081, 0.071, 0.076, 0.067, 0.071, 0.086],
    "d2_diameter":           [6, 7, 7, 7, 7, 7],
    "d3_avg_shortest_path":  [2.448, 2.470, 2.593, 2.824, 2.806, 2.715],
    "d4_avg_degree":         [2.267, 2.000, 2.133, 1.867, 2.000, 2.400],
    "d5_max_betweenness":    [0.462, 0.598, 0.505, 0.725, 0.599, 1.022],
    "d6_avg_betweenness":    [0.253, 0.197, 0.228, 0.303, 0.294, 0.307],
    "d7_max_eigenvalue":     [0.536, 0.473, 0.464, 0.354, 0.473, 0.437],
    "d8_avg_eigenvalue":     [0.195, 0.215, 0.226, 0.239, 0.215, 0.229],
})
table5.to_csv(f"{OUT_DIR}/table5_graph_descriptors.csv", index=False)

# ---------------------------------------------------------------------
# 4. Table 6 (thesis) - molar-mass-based node descriptors, 15 subdomains
#    (G1-G15) per variant.
# ---------------------------------------------------------------------
table6 = pd.DataFrame({
    "Node": [f"G{i}" for i in range(1, 16)],
    "Wildtype": [0.130105, 0.464641, 0.852632, 0.669386, 0.610677, 0.578518, 0.554695,
                 0.441309, 0.087336, 0.104682, 0.302636, 0.583973, 0.234868, 0.142609, 0.096123],
    "Y100C":    [0.143115, 0.511105, 0.803400, 0.736325, 0.537250, 0.636370, 1.018615,
                 0.248610, 0.096070, 0.115150, 0.294390, 0.852500, 0.258355, 0.156870, 0.105735],
    "C124Y":    [0.384714, 0.479168, 0.753199, 0.690315, 0.550218, 0.596606, 0.600178,
                 0.289355, 0.118207, 0.154493, 0.275995, 0.416303, 0.436483, 0.095880, 0.298481],
    "S132F":    [0.153335, 0.547603, 0.860770, 0.788906, 0.575615, 0.681813, 0.653737,
                 0.232024, 0.102930, 0.157711, 0.315412, 0.475759, 0.276804, 0.168072, 0.113285],
    "M133K":    [0.134191, 0.479236, 0.879414, 0.690413, 0.629859, 0.810863, 0.572119,
                 0.233061, 0.090080, 0.108017, 0.490206, 0.416362, 0.242246, 0.147089, 0.099142],
    "G145R":    [0.119263, 0.647712, 0.781579, 0.870992, 0.559787, 0.942446, 0.508471,
                 0.207175, 0.290258, 0.001633, 0.388533, 0.370042, 0.215296, 0.130725, 0.088113],
})
table6.to_csv(f"{OUT_DIR}/table6_molar_mass_descriptors.csv", index=False)

# ---------------------------------------------------------------------
# 5. Thesis-reported 50 ns / 200 ns mean & max RMSD summary (for cross-check
#    against values recomputed directly from the raw trajectories in
#    script 01; NOTE the thesis's own 200 ns S132F row contained a
#    transcription error, corrected in 01_stats_convergence.py by
#    recomputing directly from rmsd_200ns.csv).
# ---------------------------------------------------------------------
md_summary = pd.DataFrame({
    "Variant": ["Wildtype", "Y100C", "C124Y", "S132F", "M133K", "G145R"],
    "Mean_RMSD_50ns":  [0.457, 0.491, 0.486, 0.534, 0.547, 0.503],
    "Max_RMSD_50ns":   [0.534, 0.638, 0.611, 0.592, 0.683, 0.585],
    "Mean_RMSD_200ns": [0.697, 0.650, 0.712, 0.457, 0.850, 0.572],  # as printed in thesis (S132F row is erroneous)
    "Max_RMSD_200ns":  [0.869, 0.732, 0.824, 0.657, 0.968, 0.772],
})
md_summary.to_csv(f"{OUT_DIR}/md_summary_thesis_reported.csv", index=False)

print("Data extraction complete.")
